"""openpyxl-backed .xlsx store (Excel + LibreOffice Calc)."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from finance_tracker.model import (
    META_SHEET_NAME,
    TOTAL_LABELS,
    Category,
    CategoryType,
    DEFAULT_CATEGORIES,
    MONTH_ABBREV,
    MonthEntry,
    YEAR_COLUMN_LABEL,
    YearSheetLayout,
)
from finance_tracker.totals import build_sheet_grid, parse_amount


# Spreadsheet "dark mode" styling (Excel/Calc have no workbook dark-mode flag;
# this paints the used cells so the sheet reads as a dark theme).
DARK_BG = "1E1E1E"
DARK_HEADER_BG = "2D2D30"
DARK_TOTAL_BG = "252526"
DARK_INPUT_BG = "252526"
DARK_FG = "F3F3F3"
DARK_MUTED = "9D9D9D"
DARK_ACCENT = "4EC9B0"
DARK_BORDER = "3F3F46"


class XlsxStore:
    def __init__(self, path: str | Path) -> None:
        self.path = Path(path).expanduser().resolve()

    def describe(self) -> str:
        return str(self.path)

    def list_years(self) -> List[int]:
        if not self.path.exists():
            return []
        wb = load_workbook(self.path, data_only=False)
        years: List[int] = []
        for name in wb.sheetnames:
            if name == META_SHEET_NAME:
                continue
            if name.isdigit():
                years.append(int(name))
        wb.close()
        return sorted(years)

    def load_year(self, year: int) -> Optional[YearSheetLayout]:
        if not self.path.exists():
            return None
        wb = load_workbook(self.path, data_only=False)
        name = str(year)
        if name not in wb.sheetnames:
            wb.close()
            return None
        layout = self._read_sheet(wb[name], year)
        wb.close()
        return layout

    def ensure_xlsx_suffix(self) -> None:
        if self.path.suffix.lower() != ".xlsx":
            self.path = self.path.with_suffix(".xlsx")

    def create_workbook(
        self,
        year: int,
        categories: Optional[List[Category]] = None,
    ) -> YearSheetLayout:
        """Create the .xlsx on disk immediately with a year sheet and categories."""
        self.ensure_xlsx_suffix()
        cats = list(categories or DEFAULT_CATEGORIES)
        layout = YearSheetLayout(year=year, categories=cats, months=[], amounts={})
        layout.ensure_categories(cats)

        wb = Workbook()
        default = wb.active
        if default is not None:
            default.title = str(year)
        else:
            wb.create_sheet(str(year))

        self._write_layout(wb, layout)
        self._touch_meta(wb)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(self.path)
        wb.close()
        return layout

    def save_month(
        self,
        entry: MonthEntry,
        categories: Optional[List[Category]] = None,
    ) -> YearSheetLayout:
        self.ensure_xlsx_suffix()
        if not self.path.exists():
            self.create_workbook(entry.year, categories)

        wb = self._open_or_create()
        layout = self._prepare_year(wb, entry.year, categories)
        layout.apply_month(entry)
        self._write_layout(wb, layout)
        self._touch_meta(wb)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(self.path)
        wb.close()

        # Re-load to confirm values actually landed on disk.
        verified = self.load_year(entry.year)
        if verified is None:
            raise RuntimeError(f"Save appeared to succeed but sheet {entry.year} is missing.")
        written = verified.get_month_amounts(entry.month)
        for name, expected in entry.amounts.items():
            actual = written.get(name)
            if actual is None or float(actual) != float(expected):
                raise RuntimeError(
                    f"Value for {name} did not persist "
                    f"(expected {expected}, found {actual}) in {self.path}"
                )
        return verified

    def add_category(self, year: int, category: Category) -> YearSheetLayout:
        self.ensure_xlsx_suffix()
        if not self.path.exists():
            cats = list(DEFAULT_CATEGORIES)
            if not any(c.name == category.name for c in cats):
                cats.append(category)
            self.create_workbook(year, cats)
            layout = self.load_year(year)
            assert layout is not None
            return layout
        wb = load_workbook(self.path)
        layout = self._prepare_year(wb, year, None)
        if not any(c.name == category.name for c in layout.categories):
            layout.add_category(category)
        self._write_layout(wb, layout)
        self._touch_meta(wb)
        wb.save(self.path)
        wb.close()
        return layout

    def _open_or_create(self) -> Workbook:
        if self.path.exists():
            return load_workbook(self.path)
        wb = Workbook()
        default = wb.active
        if default is not None:
            wb.remove(default)
        return wb

    def _prepare_year(
        self,
        wb: Workbook,
        year: int,
        categories: Optional[List[Category]],
    ) -> YearSheetLayout:
        name = str(year)
        if name in wb.sheetnames:
            layout = self._read_sheet(wb[name], year)
            if categories:
                layout.ensure_categories(categories)
            return layout

        template = categories or self._template_categories(wb) or list(DEFAULT_CATEGORIES)
        layout = YearSheetLayout(year=year, categories=list(template), months=[], amounts={})
        layout.ensure_categories(template)
        wb.create_sheet(name)
        return layout

    def _template_categories(self, wb: Workbook) -> Optional[List[Category]]:
        years = sorted(int(n) for n in wb.sheetnames if n.isdigit())
        if not years:
            return None
        layout = self._read_sheet(wb[str(years[-1])], years[-1])
        return list(layout.categories)

    def _read_sheet(self, ws: Worksheet, year: int) -> YearSheetLayout:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return YearSheetLayout(year=year)

        header = [str(c) if c is not None else "" for c in rows[0]]
        months: List[str] = []
        for cell in header[2:]:
            if cell in MONTH_ABBREV:
                months.append(cell)
            elif cell == YEAR_COLUMN_LABEL:
                break

        categories: List[Category] = []
        amounts: dict = {}
        for row in rows[1:]:
            if not row or row[0] is None:
                continue
            name = str(row[0]).strip()
            if name in TOTAL_LABELS:
                break
            type_raw = str(row[1]).strip() if len(row) > 1 and row[1] is not None else "Expense"
            try:
                cat_type = CategoryType(type_raw)
            except ValueError:
                cat_type = CategoryType.EXPENSE
            categories.append(Category(name, cat_type))
            cat_amounts: dict = {}
            for i, month in enumerate(months):
                col_idx = 2 + i
                if col_idx < len(row):
                    parsed = parse_amount(row[col_idx])
                    if parsed is not None:
                        cat_amounts[month] = parsed
            amounts[name] = cat_amounts

        if not categories:
            categories = list(DEFAULT_CATEGORIES)

        return YearSheetLayout(
            year=year,
            categories=categories,
            months=months,
            amounts=amounts,
        )

    def _write_layout(self, wb: Workbook, layout: YearSheetLayout) -> None:
        name = str(layout.year)
        # Replace sheet entirely so old light-styled / stale cells cannot linger.
        if name in wb.sheetnames:
            wb.remove(wb[name])
        ws = wb.create_sheet(name, 0)

        grid = build_sheet_grid(layout)
        self._apply_dark_theme(ws, grid, layout)

        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 10
        for i in range(3, 3 + max(len(layout.months), 0) + 1):
            ws.column_dimensions[get_column_letter(i)].width = 12

        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = DARK_HEADER_BG

    def _apply_dark_theme(
        self,
        ws: Worksheet,
        grid: List[List[object]],
        layout: YearSheetLayout,
    ) -> None:
        header_fill = PatternFill("solid", fgColor=DARK_HEADER_BG)
        body_fill = PatternFill("solid", fgColor=DARK_BG)
        input_fill = PatternFill("solid", fgColor=DARK_INPUT_BG)
        total_fill = PatternFill("solid", fgColor=DARK_TOTAL_BG)
        header_font = Font(bold=True, color=DARK_ACCENT, name="Calibri", size=11)
        body_font = Font(color=DARK_FG, name="Calibri", size=11)
        muted_font = Font(color=DARK_MUTED, name="Calibri", size=11)
        total_font = Font(bold=True, color=DARK_FG, name="Calibri", size=11)
        thin = Border(
            left=Side(style="thin", color=DARK_BORDER),
            right=Side(style="thin", color=DARK_BORDER),
            top=Side(style="thin", color=DARK_BORDER),
            bottom=Side(style="thin", color=DARK_BORDER),
        )
        money_align = Alignment(horizontal="right")

        month_count = len(layout.months)
        first_month_col = 3
        year_col = first_month_col + month_count

        # Paint a dark backdrop first so unused cells aren't bright white.
        max_row = max(len(grid) + 8, 24)
        max_col = max(year_col + 2, 8)
        for r in range(1, max_row + 1):
            for c in range(1, max_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = body_fill
                cell.font = muted_font

        for r_idx, row in enumerate(grid, start=1):
            label = row[0] if row else None
            is_header = r_idx == 1
            is_total = isinstance(label, str) and label in TOTAL_LABELS

            for c_idx, value in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                cell.border = thin

                if is_header:
                    cell.fill = header_fill
                    cell.font = header_font
                elif is_total:
                    cell.fill = total_fill
                    cell.font = total_font
                elif c_idx >= first_month_col and c_idx < year_col:
                    cell.fill = input_fill
                    cell.font = body_font
                    if isinstance(value, (int, float)):
                        cell.number_format = "#,##0.00"
                        cell.alignment = money_align
                elif c_idx == year_col:
                    cell.fill = body_fill
                    cell.font = muted_font
                elif c_idx == 2:
                    cell.fill = body_fill
                    cell.font = muted_font
                else:
                    cell.fill = body_fill
                    cell.font = body_font

    def _touch_meta(self, wb: Workbook) -> None:
        if META_SHEET_NAME in wb.sheetnames:
            wb.remove(wb[META_SHEET_NAME])
        ws = wb.create_sheet(META_SHEET_NAME)
        ws["A1"] = "last_updated"
        ws["B1"] = datetime.now(timezone.utc).isoformat()
        ws["A2"] = "theme"
        ws["B2"] = "dark"
        fill = PatternFill("solid", fgColor=DARK_BG)
        font = Font(color=DARK_FG, name="Calibri")
        for row in ws.iter_rows(min_row=1, max_row=10, max_col=5):
            for cell in row:
                cell.fill = fill
                cell.font = font
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = DARK_HEADER_BG
